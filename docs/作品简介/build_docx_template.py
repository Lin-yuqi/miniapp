from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "作品介绍文档.docx"
ASSETS = ROOT / "assets"
CSU_LOGO = ASSETS / "csu-logo.png"


ACCENT = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EFF6FC"
MID_GRAY = "666666"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "7F7F7F", size: str = "8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=140, bottom=90, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_cm: list[float]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths_dxa = [int(round(w * 567)) for w in widths_cm]
    total_dxa = sum(widths_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_dxa))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for dxa in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(dxa))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, fld_text, fld_end):
        run._r.append(node)


def style_run(run, size=12, bold=False, color=None, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(p, size=12, font_name="宋体", color=None):
    for run in p.runs:
        style_run(run, size=size, color=color, font_name=font_name)


def add_center_text(doc, text, size=14, bold=False, color=None, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    style_run(r, size=size, bold=bold, color=color)
    return p


def add_body_placeholder(doc, text="【请在此处填写正文。】"):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, size=12, color=MID_GRAY)
    return p


def add_section_heading(doc, number: str, title: str):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(f"{number} {title}")
    style_run(r, size=16, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(10)
    return p


def add_subheading(doc, title: str):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(title)
    style_run(r, size=13, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_hint_box(doc, title: str, hints: list[str]):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, [14.6])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_border(cell, "A9C5DD", "8")
    set_cell_margins(cell, 120, 160, 120, 160)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    r = p.add_run(title)
    style_run(r, size=11, bold=True, color=ACCENT)
    for hint in hints:
        p = cell.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.18)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(hint)
        style_run(r, size=10.5, color=MID_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_fill_table(doc, caption: str, rows: list[tuple[str, str]], widths=(3.2, 11.4)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(caption)
    style_run(r, size=10.5, bold=True)
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_width(table, list(widths))
    for i, (label, content) in enumerate(rows):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        for cell in (c0, c1):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(c0, LIGHT_BLUE)
        if i == 0:
            set_cell_shading(c1, "EAF2F8")
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c0.paragraphs[0].add_run(label)
        style_run(c0.paragraphs[0].runs[0], size=10.5, bold=True)
        c1.paragraphs[0].add_run(content)
        style_run(c1.paragraphs[0].runs[0], size=10.5, color=MID_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    if CSU_LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(CSU_LOGO), width=Cm(2.3))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(56)
    r = p.add_run("参赛作品说明书")
    style_run(r, size=22, bold=True, font_name="黑体")

    cover_rows = [
        ("作品名称：", "WeCraft AI"),
        ("学校：", "中南大学"),
        ("学院：", "【请填写学院】"),
        ("专业班别：", "【请填写专业班级】"),
        ("队员姓名：", "【请填写队员姓名】"),
        ("指导老师：", "【请填写指导老师】"),
        ("完成时间：", "【请填写完成时间】"),
    ]
    for label, value in cover_rows:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        r1 = p.add_run(label)
        r2 = p.add_run(value)
        style_run(r1, size=14, bold=True)
        style_run(r2, size=14, bold=True if "【" not in value else False)
    doc.add_page_break()


def add_static_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录")
    style_run(r, size=18, bold=True, font_name="黑体")
    p.paragraph_format.space_after = Pt(18)
    items = [
        ("1", "简介"),
        ("2", "设计原理"),
        ("3", "队员分工"),
        ("4", "创新点"),
        ("5", "实用点"),
        ("6", "总结"),
    ]
    for num, title in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.4)
        p.paragraph_format.right_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(f"{num} {title}\t【页码】")
        style_run(r, size=12, color=MID_GRAY if "页码" in r.text else None)
    doc.add_page_break()


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.7)
    sec.right_margin = Cm(2.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in [("Heading 1", 16), ("Heading 2", 13)]:
        st = styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WeCraft AI")
    style_run(r, size=9, color=MID_GRAY)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)

    footer = section.footer
    footer.is_linked_to_previous = False
    add_page_number(footer.paragraphs[0])
    for run in footer.paragraphs[0].runs:
        style_run(run, size=9)


def add_template_sections(doc):
    add_section_heading(doc, "1", "简介")
    add_hint_box(doc, "填写提示", [
        "说明作品要解决的问题、目标用户、应用场景和整体定位。",
        "建议先用 1-2 段总述，再补充作品输入、输出和最终呈现形式。",
    ])
    add_fill_table(doc, "表 1-1 作品基本信息", [
        ("作品名称", "WeCraft AI"),
        ("作品定位", "【例如：基于多 Agent 协作的微信小程序智能生成平台】"),
        ("面向场景", "【请填写校园服务、课程设计、竞赛原型等应用场景】"),
        ("核心目标", "【请填写作品希望解决的主要问题】"),
    ])
    add_body_placeholder(doc, "【请在此处填写作品简介正文。】")

    add_section_heading(doc, "2", "设计原理")
    add_hint_box(doc, "填写提示", [
        "按流程说明系统如何从自然语言需求转化为小程序工程。",
        "可包含技术架构、主要模块、数据流、关键算法或核心机制。",
    ])
    add_fill_table(doc, "表 2-1 技术路线梳理", [
        ("需求输入", "【用户通过 Web UI 输入自然语言想法，系统进行需求澄清】"),
        ("方案整理", "【需求分析、页面结构、数据模型、权限规则、默认假设】"),
        ("工程生成", "【文件规划、页面四件套生成、组件/工具文件生成】"),
        ("质量校验", "【路径、事件、组件、云函数等一致性检查与重试】"),
    ])
    add_body_placeholder(doc, "【请在此处填写设计原理正文，可补充流程图或架构图。】")

    add_section_heading(doc, "3", "队员分工")
    add_hint_box(doc, "填写提示", [
        "按成员填写主要职责和完成内容，尽量使用可验证的任务描述。",
        "如暂未确定成员，可先保留表格占位，后续直接替换姓名和内容。",
    ])
    add_fill_table(doc, "表 3-1 队员分工表", [
        ("队员一", "【姓名：】  【负责内容：需求分析、资料整理、文档撰写等】"),
        ("队员二", "【姓名：】  【负责内容：系统设计、核心流程开发等】"),
        ("队员三", "【姓名：】  【负责内容：UI 设计、测试验证、演示材料等】"),
        ("队员四", "【姓名：】  【负责内容：代码生成策略、部署调试等】"),
    ])
    add_body_placeholder(doc, "【请在此处补充团队协作方式、版本迭代过程或个人贡献说明。】")

    add_section_heading(doc, "4", "创新点")
    add_hint_box(doc, "填写提示", [
        "突出与普通代码生成工具或传统开发流程相比的差异。",
        "建议分点说明机制创新、流程创新和工程化创新。",
    ])
    add_fill_table(doc, "表 4-1 创新点概述", [
        ("多 Agent 协作", "【说明不同 Agent 在需求、架构、上下文、代码生成中的分工】"),
        ("上下文整理", "【说明如何压缩并保留关键业务规则、页面关系和字段约束】"),
        ("一致性校验", "【说明如何减少路径、事件、组件和云函数引用错误】"),
        ("可追踪生成", "【说明中间产物留存、校验报告和生成摘要的作用】"),
    ])
    add_body_placeholder(doc, "【请在此处填写创新点详细说明。】")

    add_section_heading(doc, "5", "实用点")
    add_hint_box(doc, "填写提示", [
        "说明作品在真实学习、竞赛、课程设计或校园服务中的使用价值。",
        "可补充已有生成案例、使用成本、适用人群和推广方式。",
    ])
    add_fill_table(doc, "表 5-1 实用价值分析", [
        ("使用门槛", "【说明自然语言输入、Web UI 操作等降低门槛的设计】"),
        ("开发效率", "【说明从需求到可运行工程的时间节省】"),
        ("结果复用", "【说明输出工程如何导入微信开发者工具并二次开发】"),
        ("应用案例", "【填写校园二手书交换小程序、扫雷小游戏等案例】"),
    ])
    add_body_placeholder(doc, "【请在此处填写实用点详细说明。】")

    add_section_heading(doc, "6", "总结")
    add_hint_box(doc, "填写提示", [
        "总结作品完成情况、主要成果和不足。",
        "最后补充未来改进方向，例如增量生成、项目级自动修复、历史任务管理等。",
    ])
    add_body_placeholder(doc, "【请在此处填写总结正文。】")
    add_fill_table(doc, "表 6-1 后续改进方向", [
        ("功能完善", "【如历史任务管理、增量生成、模型切换】"),
        ("质量提升", "【如项目级自动修复、真实预览部署、自动测试】"),
        ("推广应用", "【如课程实践、校园服务项目、竞赛原型开发】"),
    ])


def build():
    doc = setup_document()
    add_cover(doc)
    add_static_toc(doc)
    # Apply running header/footer only from正文 section onward.
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    body_section = doc.sections[-1]
    body_section.top_margin = Cm(2.4)
    body_section.bottom_margin = Cm(2.2)
    body_section.left_margin = Cm(2.7)
    body_section.right_margin = Cm(2.7)
    add_header_footer(body_section)
    add_template_sections(doc)
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
